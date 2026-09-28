import sys
from decimal import Decimal as D
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from data import SANITAER, HEIZUNG, VORRECHNUNGEN, FBH_RABATT
from payments import ZAHLUNGEN, BELEGE

EUR = '#,##0.00 "€"'
B = Font(bold=True); H = PatternFill("solid", fgColor="DCE6F1"); T = PatternFill("solid", fgColor="F2F2F2")
thin = Border(bottom=Side(style="thin"))

wb = Workbook()
ws = wb.active; ws.title = "Kumulierte Aufstellung"
ws.append(["Kumulierte Schlussrechnung – BV Liesborner Weg 19 (Villa), Projekt 0200 – ENTWURF"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Auftraggeber (Empfehlung): CREST Living GmbH & Co. KG, Liesborner Weg 19, 13507 Berlin"])
ws.append([])
ws.append(["Pos.", "Beschreibung", "Menge", "Einheit", "EP netto", "GP netto", "Beleg"])
for c in ws[4]: c.font = B; c.fill = H
def add(rows):
    first = ws.max_row + 1
    for p, txt, q, u, ep, b in rows:
        if q is None:
            ws.append([p, txt]); ws.cell(ws.max_row, 2).font = Font(italic=True); continue
        r = ws.max_row + 1
        ws.append([p, txt, float(D(q)), u, float(D(ep)), f"=ROUND(C{r}*E{r},2)", b])
    return first, ws.max_row
ws.append(["1", "Sanitärobjekte liefern und montieren, bauseits gelieferte Armaturen montieren"]); ws.cell(ws.max_row, 2).font = B
s1, s2 = add(SANITAER)
ws.append(["", "Summe Titel 1 Sanitär", None, None, None, f"=SUM(F{s1}:F{s2})"]); r_san = ws.max_row
ws.append(["2", "Heizung"]); ws.cell(ws.max_row, 2).font = B
h1, h2 = add(HEIZUNG)
ws.append(["", "abzgl. Rabatt auf Pos. 2.7.0–2.11 (wie Dok. 176)", float(FBH_RABATT), "%", None, f"=-ROUND(SUM(F{h1+1}:F{h2})*C{h2+1},2)"]); r_rab = ws.max_row
ws.cell(r_rab, 3).number_format = "0%"
ws.append(["", "Summe Titel 2 Heizung", None, None, None, f"=SUM(F{h1}:F{r_rab})"]); r_hz = ws.max_row
ws.append([])
ws.append(["", "Gesamtleistung netto", None, None, None, f"=F{r_san}+F{r_hz}"]); r_gn = ws.max_row
ws.append(["", "USt 19 %", None, None, None, f"=ROUND(F{r_gn}*0.19,2)"]); r_gu = ws.max_row
ws.append(["", "Gesamtleistung brutto", None, None, None, f"=F{r_gn}+F{r_gu}"]); r_gb = ws.max_row
for r in (r_san, r_hz, r_gn, r_gb):
    for c in ws[r]: c.font = B; c.fill = T
for row in ws.iter_rows(min_row=5, min_col=5, max_col=6):
    for c in row: c.number_format = EUR
ws.column_dimensions["A"].width = 9; ws.column_dimensions["B"].width = 78
for col, w in zip("CDEFG", (9, 8, 13, 14, 12)): ws.column_dimensions[col].width = w

# Sheet 2: Rechnungen & Zahlungen
w2 = wb.create_sheet("Rechnungen & Zahlungen")
w2.append(["Alle Belege zum BV Liesborner Weg 19 (Villa)"]); w2["A1"].font = Font(bold=True, size=13)
w2.append([])
hdr = ["Beleg", "Datum", "Empfänger", "Gewerk", "netto", "USt", "brutto", "gezahlt", "offen", "zählt", "Status / Bemerkung", "Drive-Datei-ID"]
w2.append(hdr)
for c in w2[3]: c.font = B; c.fill = H
for b in BELEGE:
    r = w2.max_row + 1
    w2.append([b["nr"], b["datum"], b["an"], b["gewerk"], b["netto"], b["ust"], b["brutto"], b["gezahlt"],
               f'=IF(H{r}="","prüfen",G{r}-H{r})', "Ja" if b["zaehlt"] else "Nein", b["status"], b["id"]])
last = w2.max_row
sr = last + 1
w2.append(["Summe (nur Zeilen mit zählt = Ja)", None, None, None] +
          [f'=SUMIF($J$4:$J${last},"Ja",{c}4:{c}{last})' for c in "EFGHI"])
for c in w2[sr]: c.font = B; c.fill = T
for row in w2.iter_rows(min_row=4, min_col=5, max_col=9):
    for c in row: c.number_format = EUR
for col, w in zip("ABCDEFGHIJKL", (16, 11, 27, 32, 12, 11, 12, 12, 12, 7, 70, 36)): w2.column_dimensions[col].width = w
for row in w2.iter_rows(min_row=4, min_col=11, max_col=11):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(sys.argv[1]); print("saved", sys.argv[1])
